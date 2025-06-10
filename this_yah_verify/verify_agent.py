import argparse
import logging
import json
from datetime import datetime
from typing import List, Dict, Optional

import os

import requests
from duckduckgo_search import DDGS
import spacy
import openai
from docx import Document

nlp = spacy.load('en_core_web_sm')

logging.basicConfig(level=logging.INFO, format='%(asctime)-15s %(levelname)s %(message)s')


def duckduckgo_search(query: str, max_results: int = 5) -> List[Dict[str, str]]:
    results = []
    with DDGS() as ddgs:
        for r in ddgs.text(query, max_results=max_results):
            results.append({'title': r.get('title'), 'href': r.get('href')})
    return results


def bailii_search(query: str, max_results: int = 5) -> List[Dict[str, str]]:
    return duckduckgo_search(f"site:bailii.org {query}", max_results=max_results)


def archive_search(query: str, max_results: int = 5) -> List[Dict[str, str]]:
    url = f"https://archive.org/advancedsearch.php?q={requests.utils.quote(query)}&output=json&rows={max_results}"
    res = requests.get(url, timeout=10)
    data = res.json()
    results = []
    for doc in data.get('response', {}).get('docs', []):
        identifier = doc.get('identifier')
        title = doc.get('title')
        if identifier:
            results.append({'title': title, 'href': f"https://archive.org/details/{identifier}"})
    return results


def perform_ner(text: str) -> List[str]:
    doc = nlp(text)
    return [ent.text for ent in doc.ents]


def summarise_text(text: str, model: Optional[str] = None) -> str:
    openai.api_key = os.getenv('OPENAI_API_KEY')
    if not openai.api_key:
        logging.warning('OPENAI_API_KEY not set. Returning first 200 chars.')
        return text[:200]
    prompt = f"Summarise the following text:\n{text}"
    response = openai.ChatCompletion.create(
        model=model or 'gpt-3.5-turbo',
        messages=[{'role': 'user', 'content': prompt}],
    )
    return response.choices[0].message['content'].strip()


def cross_reference_claim(query: str, max_results: int = 5) -> List[Dict[str, str]]:
    fact_sites = ['site:fullfact.org', 'site:logicallyfacts.com', 'site:factcheck.org']
    results = []
    for site in fact_sites:
        results.extend(duckduckgo_search(f"{site} {query}", max_results=max_results))
    return results


def export_docx(report: Dict[str, any], filename: str) -> None:
    document = Document()
    document.add_heading('THIS YAH Verify Report', level=1)
    for key, value in report.items():
        document.add_heading(key, level=2)
        if isinstance(value, list):
            for item in value:
                document.add_paragraph(str(item))
        else:
            document.add_paragraph(str(value))
    document.save(filename)


def build_report(claim: str, evidence: List[Dict[str, str]], fact_checks: List[Dict[str, str]]) -> Dict[str, any]:
    entities = perform_ner(claim)
    report = {
        'Claim': claim,
        'Entities': entities,
        'Evidence Sources': evidence,
        'Fact Checks': fact_checks,
        'Timestamp': datetime.utcnow().isoformat() + 'Z'
    }
    return report


def main():
    parser = argparse.ArgumentParser(description='THIS YAH Verify assistant')
    parser.add_argument('claim', help='Claim or query to verify')
    parser.add_argument('--export-json', metavar='PATH', help='Export report to JSON file')
    parser.add_argument('--export-docx', metavar='PATH', help='Export report to DOCX file')
    args = parser.parse_args()

    logging.info('Searching open sources...')
    evidence = duckduckgo_search(args.claim)
    evidence += bailii_search(args.claim)
    evidence += archive_search(args.claim)

    logging.info('Checking fact databases...')
    fact_checks = cross_reference_claim(args.claim)

    report = build_report(args.claim, evidence, fact_checks)

    if args.export_json:
        with open(args.export_json, 'w') as f:
            json.dump(report, f, indent=2)
        logging.info(f'Report saved to {args.export_json}')

    if args.export_docx:
        export_docx(report, args.export_docx)
        logging.info(f'Report saved to {args.export_docx}')

    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
